"""Check workbook structure; optionally evaluate key formulas with `formulas`.
Requires openpyxl; formulas and numpy enable the calculation checks.
This is not an execution of Excel's GUI or local-file hyperlink permissions.
"""
from pathlib import Path
from zipfile import ZipFile
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
p=ROOT/'docs/Half_Car_Iteration_Tracker.xlsx'
w=load_workbook(p)
assert len(w.sheetnames)==9
for name in ['Run Index','Vehicle Inputs','Road and Solver','Outputs']:
 s=w[name]
 assert s['A6'].value=='RUN-001' and s['A205'].value=='RUN-200'
 assert len(s.tables)==1
 assert s.freeze_panes=='B6'
assert w['Run Index']['E6'].value=='Template / not run'
assert all(w['Outputs'].cell(r,c).value is None for r in range(6,206) for c in range(2,10))
assert len(w['Compare']._charts)==2
assert w['Bumps']['A6'].value=='RUN-001'
assert w['Bumps']['B6'].value==1 and w['Bumps']['E6'].value is None
with ZipFile(p) as z:
 assert not any('vbaProject' in name or 'externalLinks/' in name for name in z.namelist())
 assert z.testzip() is None
print('PASS: workbook package, nine sheets, 200 consistent IDs, blank outputs, template status, charts, no macros/external workbook dependencies')
try:
 import formulas,numpy as np
 from formulas.functions import sh
except ImportError:
 print('Formula-evaluation checks skipped (install formulas and numpy).')
else:
 def calc(sheet,cell,inputs):
  fn=formulas.Parser().ast(w[sheet][cell].value)[1].compile()
  args=[np.array(inputs[key],dtype=object) for key in fn.inputs]
  return np.asarray(fn(*args)).item()
 empty=sh.EMPTY
 for a,b,expected in [(empty,empty,'Not entered / incomplete'),(-10,empty,'VIOLATED - later response unreliable'),(0,20,'No negative predicted (estimate)'),(10,empty,'Not entered / incomplete'),(20,-5,'VIOLATED - later response unreliable')]:
  assert calc('Outputs','S6',{'G6':[[a]],'G6:H6':[[a,b]],'H6':[[b]]})==expected
 assert calc('Outputs','O6',{'B6':[[.012]]})==12
 assert abs(calc('Outputs','P6',{'C6':[[np.pi/180]]})-1)<1e-10
 assert calc('Outputs','O6',{'B6':[[empty]]})==''
 vals={'C6':[[420]],'C6:E6':[[420,53,76]],'D6':[[53]],'G6:H6':[[1.4,1.47]],'H6':[[1.47]]}
 assert abs(calc('Vehicle Inputs','R6',vals)-9.81*(420*1.47/2.87+53))<1e-8
 ids=['RUN-002','RUN-001']+[f'RUN-{i:03d}' for i in range(3,201)]
 names=['Second run','First run']+[empty]*198
 mapping={"'RUN INDEX'!A6:A205":np.array(ids).reshape(-1,1),"'RUN INDEX'!B6:B205":np.array(names,dtype=object).reshape(-1,1),'B5':[['RUN-001']]}
 # The third-party evaluator does not implement COUNTIFS '<>' correctly;
 # check the actual INDEX/MATCH lookup component independently. Excel's
 # COUNTIFS guard is syntax-checked below, not falsely reported as evaluated.
 lookup_formula = "=INDEX('Run Index'!B6:B205,MATCH(B5,'Run Index'!A6:A205,0))"
 fn = formulas.Parser().ast(lookup_formula)[1].compile()
 value = fn(*(np.array(mapping[key],dtype=object) for key in fn.inputs))
 assert np.asarray(value).item()=='First run'
 # Parse representative top/bottom rows, comparison fields and graph formulas.
 for s in w:
  for row in s:
   for cell in row:
    if cell.data_type=='f' and (cell.row in [6,205,2005] or s.title=='Compare'):
     formulas.Parser().ast(cell.value)
 print('PASS: contact flags for missing/negative/zero loads, SI conversions, static load, INDEX/MATCH component after reordering, and representative formula syntax (COUNTIFS nonblank guards not evaluated)')
